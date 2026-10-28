#!/usr/bin/env python3
"""Export one account's AMC results to a data workbook for the client's Drive AMC/<yyyy-mm> folder.

Reads the same folder as amc_findings.py (after it has run) and writes
<Brand>_AMC_Data_<yyyy-mm>.xlsx with one tab per report plus the SQL that produced it.

Usage: python3 tools/amc_data_workbook.py <folder> <yyyy-mm> [sql_pack.sql]
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HDR = PatternFill("solid", fgColor="111318")
F = lambda b=False, c="000000": Font(name="Arial", size=10, bold=b, color=c)


def sheet(wb, title, head, rows, widths=None, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(name="Arial", size=9, italic=True, color="666666"); r0 = 3
    for j, h in enumerate(head, 1):
        c = ws.cell(row=r0, column=j, value=h); c.font = F(True, "FFFFFF"); c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=v).font = F()
    for j, w in enumerate(widths or [14] * len(head), 1):
        ws.column_dimensions[ws.cell(row=r0, column=j).column_letter].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
    return ws


def main(d, month, sql=None):
    f = json.load(open(os.path.join(d, "amc_findings.json")))
    m, L = f["meta"], f["months"]
    wb = Workbook(); wb.remove(wb.active)
    N = f["ntb"]
    sheet(wb, "NTB Monthly", ["Month", "Ad spend", "Customers", "NTB customers", "Orders", "NTB orders", "Sales", "NTB sales", "SP NTB orders", "SB NTB orders"],
          [[L[i], N["spend"][i], N["customers"][i], N["ntb_customers"][i], N["purchases"][i], N["ntb_purchases"][i], N["sales"][i],
            N["ntb_sales"][i], N["sp_ntb"][i], N["sb_ntb"][i]] for i in range(len(L))], [12] + [13] * 9,
          "AMC ad-attributed purchases on conversion date. SD NTB is inside totals (below the privacy floor by month).")
    S = json.load(open(os.path.join(d, "search_terms.json")))
    sheet(wb, "NTB Search Terms", ["Search term", "Ad type", "NTB customers", "NTB orders", "NTB sales", "Impressions", "Clicks", "Term spend"],
          [[r["customer_search_term"] or "(below privacy floor)", r["ad_product_type"] or "", *(float(r[k] or 0) for k in
            ("ntb_customers", "ntb_purchases", "ntb_sales", "impressions", "clicks", "spend"))] for r in S],
          [44, 18, 12, 12, 12, 12, 10, 12], f"NTB shoppers who clicked the term (same campaign) within 14 days before the NTB purchase, {m['window_label']}.")
    R = f["repeat"]
    sheet(wb, "Repeat & SnS Monthly", ["Month", "Customers", "Sales", "SnS customers", "SnS sales", "Repeat customers", "Repeat sales",
                                       "One-time customers", "One-time sales", "SnS sign-ups"],
          [[L[i], R["customers"][i], R["sales"][i], R["sns_customers"][i], R["sns_sales"][i], R["repeat_customers"][i], R["repeat_sales"][i],
            R["one_customers"][i], R["one_sales"][i], R["signups"][i]] for i in range(len(L))], [12] + [13] * 9,
          "AMC conversions. Segment per customer-month: SnS subscriber > repeat (bought earlier or 2+ orders) > one-time.")
    sheet(wb, "Repeat by ASIN", ["ASIN", "Product", "Customers", "Orders", "Sales", "SnS sales", "Repeat sales", "One-time sales"],
          [[a["asin"], a["name"], a["customers"], a["orders"], a["sales"], a["sns_sales"], a["repeat_sales"], a["one_sales"]] for a in f["asin"]],
          [13, 36, 11, 10, 13, 12, 13, 13], f"{m['repeat_window']}. Customers do not sum across ASINs.")
    P = json.load(open(os.path.join(d, "path_v2.json")))
    sheet(wb, "Path to Conversion", ["Month", "Ad types seen (30 days before purchase)", "Buyers", "Orders", "Sales", "Avg touch days",
                                     "Avg days first touch to order", "Buyers who clicked"],
          [[r["month"] or "(below privacy floor)", r["path"], r["buyers"], r["orders"], r["sales"], r["avg_touch_days"],
            r["avg_days_to_order"], r["buyers_clicked"]] for r in sorted(P, key=lambda r: (r["month"] or "9999", -r["buyers"]))],
          [16, 52, 10, 10, 12, 12, 14, 12])
    sheet(wb, "Time to Conversion", ["Ad type", "Click to purchase", "Purchases", "NTB purchases", "Sales"],
          [[t["ad"], t["bucket"], t["purchases"], t["ntb"], t["sales"]] for t in f["ttc"]], [20, 16, 12, 12, 12],
          f"Click-attributed purchases, {m['window_label']}.")
    if sql and os.path.exists(sql):
        ws = wb.create_sheet("AMC SQL"); ws.column_dimensions["A"].width = 140
        for i, line in enumerate(open(sql).read().splitlines(), 1):
            ws.cell(row=i, column=1, value=line).font = Font(name="Courier New", size=9)
    ws = wb.create_sheet("About", 0); ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 110
    rows = [("Account", m["brand"]), ("Marketplace", m["marketplace"]), ("Run", month)] + [tuple(w) for w in m["windows"]] + [("Note", n) for n in m.get("notes", [])]
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(row=i, column=1, value=a).font = F(True); ws.cell(row=i, column=2, value=b).font = F()
        ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True)
    out = os.path.join(d, f"{m['brand'].replace(' ', '_')}_AMC_Data_{month}.xlsx")
    wb.save(out); print("[done]", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
